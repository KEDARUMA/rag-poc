"""RAG検索テスト画面のローカルWebサーバー。"""

import argparse
import hashlib
import json
import logging
import mimetypes
import sys
import subprocess
import tempfile
import threading
import time
import uuid
from email import policy
from email.parser import BytesParser
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote

# 定数: src配下のRAG実装を読み込むためのパス。
PROJECT_DIRECTORY = Path(__file__).resolve().parent
sys.path.insert(0, str(PROJECT_DIRECTORY / "src"))

from rag_poc.config import DATA_ROOT, TEMP_ROOT, TEXT_CHUNK_OVERLAP, TEXT_CHUNK_SIZE
from rag_poc.indexer import index_directory
from rag_poc.searcher import search_many

# 定数: 検索APIの診断ログ。
LOGGER = logging.getLogger(__name__)

# 定数: 画面から検索クエリとして受け付ける拡張子。
ALLOWED_QUERY_EXTENSIONS = {
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
    ".mp4",
    ".mov",
    ".mkv",
    ".avi",
    ".webm",
    ".pdf",
    ".xlsx",
}
# 定数: 検索結果プレビューで配信する登録対象の拡張子。
ALLOWED_PREVIEW_EXTENSIONS = {
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
    ".mp4",
    ".mov",
    ".mkv",
    ".avi",
    ".webm",
}
# 定数: ブラウザー再生用に変換する動画拡張子。
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".avi", ".webm"}
# 定数: ブラウザへ返す静的ファイルとMIMEタイプ。
STATIC_FILES = {
    "/": ("index.html", "text/html; charset=utf-8"),
    "/ui.css": ("ui.css", "text/css; charset=utf-8"),
    "/ui.js": ("ui.js", "text/javascript; charset=utf-8"),
}
# 状態: インデックス更新ジョブの共有状態。
INDEX_JOB_STATE = {
    "status": "idle",
    "processed": 0,
    "total": 0,
    "current_file": "",
    "logs": [],
    "summary": None,
    "error": "",
}
# 状態: インデックス更新ジョブの共有状態を保護するロック。
INDEX_JOB_LOCK = threading.Lock()


def _record_index_progress(progress: dict) -> None:
    """インデクサの進捗とログを共有状態へ反映する。"""
    event = progress.get("event")
    message = str(progress.get("message", ""))
    with INDEX_JOB_LOCK:
        if event == "started":
            INDEX_JOB_STATE["processed"] = 0
            INDEX_JOB_STATE["total"] = int(progress.get("total", 0))
        elif event == "file_started":
            INDEX_JOB_STATE["current_file"] = Path(
                str(progress.get("path", ""))
            ).name
        elif event == "file_completed":
            INDEX_JOB_STATE["processed"] = int(progress.get("processed", 0))
            INDEX_JOB_STATE["total"] = int(progress.get("total", 0))
            INDEX_JOB_STATE["current_file"] = ""

        if message:
            level = "error" if progress.get("result") == "error" else "info"
            INDEX_JOB_STATE["logs"].append(
                {"message": message, "level": level}
            )


def _run_index_job() -> None:
    """既存インデックスを消去して全件再登録する。"""
    try:
        summary = index_directory(
            root=DATA_ROOT,
            rebuild=True,
            progress_callback=_record_index_progress,
        )
    except Exception as error:
        with INDEX_JOB_LOCK:
            INDEX_JOB_STATE["status"] = "failed"
            INDEX_JOB_STATE["current_file"] = ""
            INDEX_JOB_STATE["error"] = str(error)
            INDEX_JOB_STATE["logs"].append(
                {"message": f"更新に失敗しました: {error}", "level": "error"}
            )
        return

    error_count = len(summary.get("errors", []))
    if error_count:
        result_message = (
            f"更新完了（一部エラー）: 登録 {summary['indexed']}件、"
            f"スキップ {summary['skipped']}件、エラー {error_count}件"
        )
    else:
        result_message = (
            f"更新完了: 登録 {summary['indexed']}件、"
            f"スキップ {summary['skipped']}件"
        )
    with INDEX_JOB_LOCK:
        INDEX_JOB_STATE["status"] = (
            "completed_with_errors" if error_count else "completed"
        )
        INDEX_JOB_STATE["current_file"] = ""
        INDEX_JOB_STATE["summary"] = summary
        INDEX_JOB_STATE["logs"].append(
            {
                "message": result_message,
                "level": "error" if error_count else "info",
            }
        )


def _extract_text_query(path: Path) -> str:
    """TXT・PDF・XLSXから検索用の文字列を抽出する。"""
    suffix = path.suffix.casefold()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8-sig", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(filename=path, read_only=True, data_only=True)
        try:
            rows = []
            for worksheet in workbook.worksheets:
                rows.append(worksheet.title)
                for row in worksheet.iter_rows(values_only=True):
                    values = [
                        str(value)
                        for value in row
                        if value is not None and str(value).strip()
                    ]
                    if values:
                        rows.append("\t".join(values))
            return "\n".join(rows)
        finally:
            workbook.close()
    raise ValueError(f"文字抽出に未対応の形式です: {path.name}")


def _split_text_query(text: str) -> list[str]:
    """長い文書クエリをEmbedding向けの重なり付きチャンクへ分割する。"""
    normalized_text = text.strip()
    if not normalized_text:
        return []
    if len(normalized_text) <= TEXT_CHUNK_SIZE:
        return [normalized_text]
    step = TEXT_CHUNK_SIZE - TEXT_CHUNK_OVERLAP
    return [
        normalized_text[start : start + TEXT_CHUNK_SIZE]
        for start in range(0, len(normalized_text), step)
    ]


def _filename(path: str) -> str:
    """Windows形式とLinux形式のパスから表示用ファイル名を取り出す。"""
    return path.replace("\\", "/").rsplit("/", 1)[-1]


def _create_webm_preview(source_path: Path) -> Path:
    """元動画をChromium再生用のVP9 WebMへ変換してキャッシュする。"""
    source_stat = source_path.stat()
    cache_root = TEMP_ROOT / "webm_preview_cache"
    cache_root.mkdir(parents=True, exist_ok=True)
    cache_identity = (
        f"{source_path}:{source_stat.st_size}:{source_stat.st_mtime_ns}"
    ).encode("utf-8")
    cache_key = hashlib.sha256(cache_identity).hexdigest()
    cached_path = cache_root / f"{cache_key}.webm"
    if cached_path.is_file() and cached_path.stat().st_size > 0:
        return cached_path

    temporary_path = cache_root / f"{cache_key}.{uuid.uuid4().hex}.tmp.webm"
    command = [
        "ffmpeg",
        "-hide_banner",
        "-loglevel",
        "error",
        "-y",
        "-i",
        str(source_path),
        "-map",
        "0:v:0",
        "-map",
        "0:a:0?",
        "-vf",
        "scale=w='min(1920,iw)':h=-2",
        "-c:v",
        "libvpx-vp9",
        "-deadline",
        "realtime",
        "-cpu-used",
        "8",
        "-crf",
        "32",
        "-b:v",
        "0",
        "-c:a",
        "libopus",
        "-b:a",
        "96k",
        "-f",
        "webm",
        str(temporary_path),
    ]
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            check=False,
        )
        if completed.returncode != 0 or not temporary_path.is_file():
            detail = completed.stderr.strip() or "変換後ファイルがありません"
            raise RuntimeError(f"動画をWebMへ変換できませんでした: {detail}")
        temporary_path.replace(cached_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return cached_path


class RAGRequestHandler(BaseHTTPRequestHandler):
    """画面の静的ファイルと検索リクエストを処理する。"""

    def do_GET(self) -> None:
        """画面と静的ファイルを返す。"""
        request_path, _, query_string = self.path.partition("?")
        if request_path == "/api/index/status":
            with INDEX_JOB_LOCK:
                payload = dict(INDEX_JOB_STATE)
                payload["logs"] = [dict(log) for log in INDEX_JOB_STATE["logs"]]
            self._send_json(200, payload)
            return
        if request_path == "/api/file":
            query = parse_qs(query_string)
            file_path = query.get("path", [""])[0]
            preview_format = query.get("format", [""])[0]
            self._serve_data_file(file_path, preview_format)
            return

        file_info = STATIC_FILES.get(request_path)
        if file_info is None:
            self.send_error(404)
            return
        file_name, content_type = file_info
        content = (PROJECT_DIRECTORY / "ui" / file_name).read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(content)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(content)

    def _serve_data_file(
        self, path_value: str, preview_format: str = ""
    ) -> None:
        """rag/data配下の検索結果ファイルをインライン配信する。"""
        try:
            data_root = DATA_ROOT.resolve(strict=True)
            file_path = Path(path_value).resolve(strict=True)
            file_path.relative_to(data_root)
        except (OSError, RuntimeError, ValueError):
            self.send_error(404)
            return

        if (
            not file_path.is_file()
            or file_path.suffix.casefold() not in ALLOWED_PREVIEW_EXTENSIONS
        ):
            self.send_error(404)
            return

        if preview_format not in {"", "webm"}:
            self.send_error(400)
            return
        response_name = file_path.name
        if preview_format == "webm":
            if file_path.suffix.casefold() not in VIDEO_EXTENSIONS:
                self.send_error(400)
                return
            try:
                file_path = _create_webm_preview(file_path)
            except (OSError, RuntimeError) as error:
                self.send_error(500, str(error))
                return
            response_name = Path(path_value).name.rsplit(".", 1)[0] + ".webm"

        file_size = file_path.stat().st_size
        start = 0
        end = file_size - 1
        range_header = self.headers.get("Range")
        is_partial = bool(range_header)
        if range_header:
            try:
                unit, range_value = range_header.split("=", 1)
                if unit != "bytes" or "," in range_value:
                    raise ValueError("未対応のRange指定です")
                first_byte, last_byte = range_value.split("-", 1)
                if first_byte:
                    start = int(first_byte)
                    end = int(last_byte) if last_byte else file_size - 1
                else:
                    suffix_length = int(last_byte)
                    if suffix_length < 1:
                        raise ValueError("未対応のRange指定です")
                    start = max(file_size - suffix_length, 0)
                    end = file_size - 1
                if start < 0 or start >= file_size or end < start:
                    raise ValueError("Rangeがファイル範囲外です")
                end = min(end, file_size - 1)
            except ValueError:
                self.send_response(416)
                self.send_header("Content-Range", f"bytes */{file_size}")
                self.send_header("Content-Length", "0")
                self.end_headers()
                return

        content_length = max(0, end - start + 1)
        content_type, _ = mimetypes.guess_type(file_path.name)
        if preview_format == "webm":
            content_type = "video/webm"
        if file_path.suffix.casefold() == ".txt":
            content_type = "text/plain; charset=utf-8"
        self.send_response(206 if is_partial else 200)
        self.send_header("Content-Type", content_type or "application/octet-stream")
        self.send_header("Content-Length", str(content_length))
        self.send_header(
            "Content-Disposition",
            f"inline; filename*=UTF-8''{quote(response_name)}",
        )
        self.send_header("Accept-Ranges", "bytes")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Cache-Control", "no-store")
        if is_partial:
            self.send_header("Content-Range", f"bytes {start}-{end}/{file_size}")
        self.end_headers()

        with file_path.open("rb") as file_stream:
            file_stream.seek(start)
            remaining = content_length
            while remaining > 0:
                chunk = file_stream.read(min(64 * 1024, remaining))
                if not chunk:
                    break
                self.wfile.write(chunk)
                remaining -= len(chunk)

    def do_POST(self) -> None:
        """インデックス更新または検索を実行し、結果をJSONで返す。"""
        if self.path == "/api/index":
            self._start_index_job()
            return
        if self.path != "/api/search":
            self.send_error(404)
            return
        request_id = uuid.uuid4().hex[:12]
        request_started = time.perf_counter()
        LOGGER.info(
            "search.request.started id=%s content_length=%s",
            request_id,
            self.headers.get("Content-Length", "unknown"),
        )
        with INDEX_JOB_LOCK:
            index_is_running = INDEX_JOB_STATE["status"] == "running"
        if index_is_running:
            LOGGER.warning(
                "search.request.rejected id=%s reason=index_running", request_id
            )
            self._send_json(
                409,
                {"error": "インデックス更新中は検索できません。完了後に再試行してください。"},
            )
            return
        try:
            query_text, uploaded_files = self._read_multipart()
            LOGGER.info(
                "search.multipart.parsed id=%s query_characters=%d file_count=%d",
                request_id,
                len(query_text),
                len(uploaded_files),
            )
            with tempfile.TemporaryDirectory(
                prefix="rag-ui-query-",
                dir=TEMP_ROOT,
            ) as temporary_directory:
                queries = []
                queries.extend(
                    ("text", chunk) for chunk in _split_text_query(query_text)
                )
                for file_index, (file_name, content) in enumerate(
                    uploaded_files, start=1
                ):
                    suffix = Path(file_name.replace("\\", "/")).suffix.casefold()
                    if suffix not in ALLOWED_QUERY_EXTENSIONS:
                        raise ValueError(f"未対応のファイル形式です: {file_name}")
                    temporary_path = (
                        Path(temporary_directory)
                        / f"{uuid.uuid4().hex}{suffix}"
                    )
                    temporary_path.write_bytes(content)
                    LOGGER.info(
                        "search.upload.saved id=%s file_index=%d extension=%s size_bytes=%d",
                        request_id,
                        file_index,
                        suffix,
                        len(content),
                    )
                    if suffix in {".txt", ".pdf", ".xlsx"}:
                        extracted_text = _extract_text_query(temporary_path)
                        chunks = _split_text_query(extracted_text)
                        if not chunks:
                            raise ValueError(
                                f"検索語として抽出できる文字がありません: {file_name}"
                            )
                        queries.extend(("text", chunk) for chunk in chunks)
                    elif suffix in {".jpg", ".jpeg", ".png"}:
                        queries.append(("image", str(temporary_path)))
                    else:
                        queries.append(("video", str(temporary_path)))

                LOGGER.info(
                    "search.queries.ready id=%s query_count=%d query_kinds=%s",
                    request_id,
                    len(queries),
                    ",".join(query_kind for query_kind, _ in queries),
                )
                LOGGER.info("search.execute.started id=%s", request_id)
                search_started = time.perf_counter()
                search_result = search_many(queries, limit=10)
                LOGGER.info(
                    "search.execute.completed id=%s result_count=%d elapsed_seconds=%.3f",
                    request_id,
                    len(search_result["results"]),
                    time.perf_counter() - search_started,
                )
                results = [
                    {
                        "name": _filename(result["path"]),
                        "path": result["path"],
                        "type": result["type"],
                        "score": result["score"],
                        "time": result.get("time"),
                        "start": result.get("start"),
                        "end": result.get("end"),
                        "width": result.get("width"),
                        "height": result.get("height"),
                        "chunk_index": result.get("chunk_index"),
                    }
                    for result in search_result["results"]
                    if result.get("path")
                ]
            self._send_json(200, {"results": results})
            LOGGER.info(
                "search.request.completed id=%s status=200 elapsed_seconds=%.3f",
                request_id,
                time.perf_counter() - request_started,
            )
        except ValueError as error:
            LOGGER.warning(
                "search.request.failed id=%s status=400 error_type=%s error=%s elapsed_seconds=%.3f",
                request_id,
                type(error).__name__,
                error,
                time.perf_counter() - request_started,
            )
            self._send_json(400, {"error": str(error)})
        except Exception as error:
            LOGGER.exception(
                "search.request.failed id=%s status=500 error_type=%s elapsed_seconds=%.3f",
                request_id,
                type(error).__name__,
                time.perf_counter() - request_started,
            )
            self._send_json(500, {"error": str(error)})

    def _start_index_job(self) -> None:
        """インデックス再構築ジョブを開始する。"""
        with INDEX_JOB_LOCK:
            if INDEX_JOB_STATE["status"] == "running":
                already_running = True
            else:
                already_running = False
                INDEX_JOB_STATE.update(
                    {
                        "status": "running",
                        "processed": 0,
                        "total": 0,
                        "current_file": "",
                        "logs": [
                            {"message": "インデックス再構築を開始します。", "level": "info"}
                        ],
                        "summary": None,
                        "error": "",
                    }
                )

        if already_running:
            self._send_json(
                409,
                {"error": "インデックス更新はすでに実行中です。"},
            )
            return

        worker = threading.Thread(target=_run_index_job, daemon=True)
        try:
            worker.start()
        except Exception as error:
            with INDEX_JOB_LOCK:
                INDEX_JOB_STATE["status"] = "failed"
                INDEX_JOB_STATE["error"] = str(error)
                INDEX_JOB_STATE["logs"].append(
                    {"message": f"更新を開始できませんでした: {error}", "level": "error"}
                )
            self._send_json(500, {"error": str(error)})
            return
        self._send_json(202, {"status": "running"})

    def _read_multipart(self) -> tuple[str, list[tuple[str, bytes]]]:
        """ブラウザのmultipartリクエストからテキストとファイルを読む。"""
        content_type = self.headers.get("Content-Type", "")
        content_length = int(self.headers.get("Content-Length", "0"))
        if not content_type.startswith("multipart/form-data") or content_length < 1:
            raise ValueError("検索入力を読み取れません")
        message = BytesParser(policy=policy.default).parsebytes(
            (
                f"Content-Type: {content_type}\r\n"
                "MIME-Version: 1.0\r\n\r\n"
            ).encode("ascii")
            + self.rfile.read(content_length)
        )
        if not message.is_multipart():
            raise ValueError("ファイル入力の形式が正しくありません")

        query_text = ""
        uploaded_files = []
        for part in message.iter_parts():
            field_name = part.get_param("name", header="content-disposition")
            file_name = part.get_filename()
            content = part.get_payload(decode=True) or b""
            if field_name == "query":
                charset = part.get_content_charset() or "utf-8"
                query_text = content.decode(charset, errors="replace")
            elif field_name == "files" and file_name:
                uploaded_files.append((file_name, content))
        return query_text, uploaded_files

    def _send_json(self, status: int, payload: dict) -> None:
        """UTF-8のJSONレスポンスを返す。"""
        content = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(content)


def main() -> int:
    """ローカルのRAG検索画面サーバーを起動する。"""
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(name)s %(message)s",
    )
    parser = argparse.ArgumentParser(description="RAG検索テスト画面")
    parser.add_argument("--host", default="127.0.0.1", help="待ち受けアドレス")
    parser.add_argument("--port", type=int, default=8765, help="待ち受けポート")
    arguments = parser.parse_args()

    server = HTTPServer((arguments.host, arguments.port), RAGRequestHandler)
    print(f"RAG検索テスト画面: http://localhost:{arguments.port}/")
    server.serve_forever()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
